from django.test import TestCase
from datetime import date
import uuid

from django.contrib.gis.geos import Point, Polygon

from types import SimpleNamespace

from django.test import SimpleTestCase
from django.test import override_settings
from pathlib import Path
from tempfile import TemporaryDirectory

from .views import _beneficiary_qs, _depense_qs, _filter_projects
from apps.accounts.models import User
from apps.organizations.models import Organization
from apps.projects.models import Project
from apps.campaigns.models import Campaign
from apps.beneficiaries.models import Beneficiary
from apps.zones.models import Department, Region, Zone
from apps.finance.models import Budget, Depense, Don
from rest_framework.test import APIClient
from rest_framework_simplejwt.tokens import RefreshToken

# Create your tests here.


class ReportOrganizationScopeTests(SimpleTestCase):
    def test_reports_return_empty_querysets_without_an_organization(self):
        user = SimpleNamespace(organization_id=None, organization=None, role="GERANT")

        for queryset in (
            _filter_projects(user),
            _beneficiary_qs(user, {}),
            _depense_qs(user, {}),
        ):
            with self.subTest(model=queryset.model.__name__):
                self.assertFalse(queryset.exists())


class OrganizationProjectAPIScopeTests(TestCase):
    def setUp(self):
        self.org_a = Organization.objects.create(
            name="ONG A", email="a@example.test", phone="770000001",
            address="Dakar", region="Dakar", country="Senegal",
            intervention_domain="Humanitaire",
        )
        self.org_b = Organization.objects.create(
            name="ONG B", email="b@example.test", phone="770000002",
            address="Thiès", region="Thiès", country="Senegal",
            intervention_domain="Humanitaire",
        )
        self.manager_a = self._user("manager-a@example.test", User.Role.GERANT, self.org_a)
        self.chef_a = self._user("chef-a@example.test", User.Role.CHEF_PROJET, self.org_a)
        self.finance_a = self._user("finance-a@example.test", User.Role.FINANCE, self.org_a)
        self.manager_b = self._user("manager-b@example.test", User.Role.GERANT, self.org_b)
        self.chef_b = self._user("chef-b@example.test", User.Role.CHEF_PROJET, self.org_b)
        self.finance_b = self._user("finance-b@example.test", User.Role.FINANCE, self.org_b)
        self.project_a = self._project("Projet A", self.org_a, self.manager_a, self.chef_a, self.finance_a)
        self.project_b = self._project("Projet B", self.org_b, self.manager_b, self.chef_b, self.finance_b)
        self.client = APIClient()
        self.client.force_authenticate(self.manager_a)

    @staticmethod
    def _user(email, role, organization):
        return User.objects.create_user(
            email=email, password="SafePass123!", role=role,
            status=User.Status.ACTIVE, is_active=True,
            first_name="Test", last_name="User", organization=organization,
        )

    @staticmethod
    def _project(name, organization, creator, chef, finance):
        return Project.objects.create(
            name=name, description="", region="Dakar", objectif="Aider les familles",
            start_date=date(2026, 1, 1), end_date=date(2026, 12, 31),
            chef_projet=chef, responsable_finance=finance,
            organization=organization, created_by=creator,
        )

    def test_project_report_lists_only_the_requesters_organization(self):
        response = self.client.get("/api/reports/projets/")

        self.assertEqual(response.status_code, 200)
        project_names = [item["nom"] for item in response.data["projets"]]
        self.assertIn(self.project_a.name, project_names)
        self.assertNotIn(self.project_b.name, project_names)

    def test_project_reports_are_available_symmetrically_to_both_tenants(self):
        for manager, own, foreign in (
            (self.manager_a, self.project_a, self.project_b),
            (self.manager_b, self.project_b, self.project_a),
        ):
            with self.subTest(organization=manager.organization_id):
                self.client.force_authenticate(manager)
                response = self.client.get("/api/reports/projets/")
                self.assertEqual(response.status_code, 200)
                names = [row["nom"] for row in response.data["projets"]]
                self.assertIn(own.name, names)
                self.assertNotIn(foreign.name, names)

    def test_report_query_parameters_cannot_select_another_organization(self):
        campaign_a = Campaign.objects.create(
            projet=self.project_a, organization=self.org_a,
            nom="Campaign A query scope", code_campagne="QUERY-A-001",
            description="Query parameter isolation", date_debut=date(2026, 1, 1),
            date_fin=date(2026, 12, 31), created_by=self.manager_a,
        )
        campaign_b = Campaign.objects.create(
            projet=self.project_b, organization=self.org_b,
            nom="Campaign B query scope", code_campagne="QUERY-B-001",
            description="Query parameter isolation", date_debut=date(2026, 1, 1),
            date_fin=date(2026, 12, 31), created_by=self.manager_b,
        )

        # Tenant A cannot use project/campaign/organization query parameters
        # to select tenant B's report data.
        response = self.client.get(
            "/api/reports/projets/",
            {"organization": self.org_b.pk, "project_id": self.project_b.pk},
        )
        self.assertEqual(response.status_code, 200)
        names = [item["nom"] for item in response.data["projets"]]
        self.assertIn(self.project_a.name, names)
        self.assertNotIn(self.project_b.name, names)

        response = self.client.get(
            "/api/reports/campagnes/", {"campaign_id": campaign_b.pk},
        )
        self.assertEqual(response.status_code, 404)

        # The same foreign selection is refused in the reverse direction.
        self.client.force_authenticate(self.manager_b)
        response = self.client.get(
            "/api/reports/campagnes/", {"campaign_id": campaign_a.pk},
        )
        self.assertEqual(response.status_code, 404)

        own_response = self.client.get(
            "/api/reports/campagnes/", {"campaign_id": campaign_b.pk},
        )
        self.assertEqual(own_response.status_code, 200)
        self.assertEqual(
            [item["id"] for item in own_response.data["campagnes"]],
            [str(campaign_b.pk)],
        )

    def test_suspended_manager_cannot_read_project_reports(self):
        access_token = str(RefreshToken.for_user(self.manager_a).access_token)
        self.manager_a.status = User.Status.SUSPENDED
        self.manager_a.save(update_fields=["status"])
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {access_token}")

        response = self.client.get("/api/reports/projets/")

        self.assertEqual(response.status_code, 403)

    def test_suspended_manager_access_to_other_report_routes(self):
        paths = (
            "/api/reports/dashboard/",
            "/api/reports/campagnes/",
            "/api/reports/beneficiaires/",
            "/api/reports/finance/",
            "/api/reports/filters/",
            "/api/reports/export/pdf/",
            "/api/reports/export/excel/",
        )
        access_token = str(RefreshToken.for_user(self.manager_a).access_token)
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {access_token}")

        active_statuses = {
            path: self.client.get(path).status_code for path in paths
        }
        self.assertTrue(
            all(code == 200 for code in active_statuses.values()),
            active_statuses,
        )

        self.manager_a.status = User.Status.SUSPENDED
        self.manager_a.save(update_fields=["status"])
        suspended_statuses = {
            path: self.client.get(path).status_code for path in paths
        }
        self.assertTrue(
            all(code == 403 for code in suspended_statuses.values()),
            suspended_statuses,
        )

    def test_foreign_project_cannot_be_read_patched_or_archived(self):
        url = f"/api/projects/{self.project_b.pk}/"

        self.assertEqual(self.client.get(url).status_code, 404)
        self.assertEqual(self.client.patch(url, {"name": "Intrusion"}, format="json").status_code, 404)
        self.assertEqual(self.client.delete(url).status_code, 404)
        self.project_b.refresh_from_db()
        self.assertEqual(self.project_b.name, "Projet B")
        self.assertFalse(self.project_b.archived)

    def test_project_creation_rejects_foreign_organization_and_assignees(self):
        payload = {
            "name": "Nouveau projet", "description": "", "region": "Dakar",
            "objectif": "Aider les familles", "start_date": "2026-01-01",
            "end_date": "2026-12-31", "chef_projet": self.chef_b.pk,
            "responsable_finance": self.finance_b.pk,
            "criteres": [{"nom": "Vulnérabilité", "poids": 100}],
            "organization_id": self.org_b.pk,
        }

        response = self.client.post("/api/projects/", payload, format="json")

        self.assertEqual(response.status_code, 400)
        self.assertFalse(Project.objects.filter(name="Nouveau projet").exists())

    def test_campaign_report_excludes_foreign_organization_beneficiaries(self):
        """Regression probe: report aggregates must defend against mismatched FK data."""
        campaign = Campaign.objects.create(
            projet=self.project_a,
            organization=self.org_a,
            nom="Campagne A",
            code_campagne="REPORT-A-001",
            description="Test report scope",
            date_debut=date(2026, 1, 1),
            date_fin=date(2026, 12, 31),
            created_by=self.manager_a,
        )
        geometry = Polygon((
            (-18.0, 14.0), (-16.0, 14.0), (-16.0, 16.0),
            (-18.0, 16.0), (-18.0, 14.0),
        ), srid=4326)
        region = Region.objects.create(nom="Report Scope Region", geometrie=geometry)
        department = Department.objects.create(
            nom="Report Scope Department", region=region, geometrie=geometry,
        )
        foreign_zone = Zone.objects.create(
            nom="Zone ONG B", region=region, departement=department,
            latitude=15, longitude=-17,
            location=Point(-17, 15, srid=4326), rayon=1,
            organization=self.org_b, created_by=self.manager_b,
        )
        # Deliberately inconsistent links reproduce data that can enter through
        # ORM/admin/import paths outside the validated API serializers.
        campaign.zones.add(foreign_zone)
        Beneficiary.objects.create(
            campagne=campaign, zone=foreign_zone, organization=self.org_b,
            nom="Ba", prenom="Awa", telephone="770000999", sexe="F",
            date_naissance=date(1990, 1, 1), latitude=15, longitude=-17,
            local_id=uuid.uuid4(), device_id="audit", created_by=self.manager_b,
        )
        own_zone = Zone.objects.create(
            nom="Zone ONG A", region=region, departement=department,
            latitude=15, longitude=-17,
            location=Point(-17, 15, srid=4326), rayon=1,
            organization=self.org_a, created_by=self.manager_a,
        )
        campaign.zones.add(own_zone)
        Beneficiary.objects.create(
            campagne=campaign, zone=own_zone, organization=self.org_a,
            nom="Diallo", prenom="Aminata", telephone="770000998", sexe="F",
            date_naissance=date(1990, 1, 1), latitude=15, longitude=-17,
            local_id=uuid.uuid4(), device_id="audit", created_by=self.manager_a,
        )

        response = self.client.get("/api/reports/campagnes/")

        self.assertEqual(response.status_code, 200)
        row = next(item for item in response.data["campagnes"] if item["id"] == str(campaign.pk))
        self.assertEqual(row["nb_beneficiaires"], 1)
        self.assertEqual(row["nb_femmes"], 1)
        self.assertEqual(row["nb_zones"], 1)

        self.client.force_authenticate(self.manager_b)
        foreign_response = self.client.get(
            "/api/reports/campagnes/", {"campaign_id": campaign.pk},
        )
        self.assertEqual(foreign_response.status_code, 404)

        self.client.force_authenticate(self.manager_a)
        excel_response = self.client.get("/api/reports/export/excel/")
        self.assertEqual(excel_response.status_code, 200)
        import openpyxl
        from io import BytesIO
        workbook = openpyxl.load_workbook(BytesIO(excel_response.content), read_only=True)
        campaign_sheet = workbook["Campagnes"]
        exported = next(row for row in campaign_sheet.iter_rows(min_row=2, values_only=True)
                        if row[0] == "REPORT-A-001")
        self.assertEqual(exported[6:10], (1, 1, 0, 1))
        self.client.force_authenticate(self.manager_b)
        foreign_excel = self.client.get(
            "/api/reports/export/excel/", {"project_id": self.project_a.pk},
        )
        self.assertEqual(foreign_excel.status_code, 404)

    def test_reports_exclude_foreign_and_inconsistent_relationship_chains(self):
        campaign_a = Campaign.objects.create(
            projet=self.project_a, organization=self.org_a,
            nom="Campagne lineage A", code_campagne="LINEAGE-A-001",
            description="", date_debut=date(2026, 1, 1), date_fin=date(2026, 12, 31),
            created_by=self.manager_a,
        )
        campaign_b = Campaign.objects.create(
            projet=self.project_b, organization=self.org_b,
            nom="Campagne lineage B", code_campagne="LINEAGE-B-001",
            description="", date_debut=date(2026, 1, 1), date_fin=date(2026, 12, 31),
            created_by=self.manager_b,
        )
        polygon = Polygon(((-18, 14), (-16, 14), (-16, 16), (-18, 16), (-18, 14)), srid=4326)
        region = Region.objects.create(nom="Lineage Region", geometrie=polygon)
        department = Department.objects.create(
            nom="Lineage Department", region=region, geometrie=polygon,
        )

        def make_zone(name, organization, creator):
            return Zone.objects.create(
                nom=name, region=region, departement=department,
                latitude=15, longitude=-17, location=Point(-17, 15, srid=4326),
                rayon=100, organization=organization, created_by=creator,
            )

        zone_a = make_zone("Lineage Zone A", self.org_a, self.manager_a)
        zone_b = make_zone("Lineage Zone B", self.org_b, self.manager_b)
        campaign_a.zones.add(zone_a, zone_b)  # Corrupt cross-tenant M2M relation.

        def make_beneficiary(name, organization, campaign, zone, creator):
            return Beneficiary.objects.create(
                campagne=campaign, zone=zone, organization=organization,
                nom=name, prenom="Audit", telephone=f"77{uuid.uuid4().int % 10**7:07d}",
                sexe="F", date_naissance=date(1990, 1, 1), latitude=15, longitude=-17,
                local_id=uuid.uuid4(), device_id="report-audit", created_by=creator,
                score_vulnerabilite=88,
            )

        own_beneficiary = make_beneficiary("Beneficiary A", self.org_a, campaign_a, zone_a, self.manager_a)
        make_beneficiary("Beneficiary B", self.org_b, campaign_b, zone_b, self.manager_b)
        make_beneficiary("Foreign campaign link", self.org_a, campaign_b, zone_b, self.manager_a)
        make_beneficiary("Foreign zone link", self.org_a, campaign_a, zone_b, self.manager_a)

        budget_a = Budget.objects.create(
            organization=self.org_a, projet=self.project_a, montant=100,
            source_financement="Audit", created_by=self.manager_a,
        )
        Budget.objects.create(
            organization=self.org_b, projet=self.project_b, montant=1000,
            source_financement="Foreign", created_by=self.manager_b,
        )
        Budget.objects.create(
            organization=self.org_a, projet=self.project_b, montant=700,
            source_financement="Corrupt foreign project", created_by=self.manager_a,
        )
        Depense.objects.create(
            reference="LINEAGE-EXP-A", organization=self.org_a, projet=self.project_a,
            budget=budget_a, categorie="A", montant=10, created_by=self.manager_a,
        )
        Depense.objects.create(
            reference="LINEAGE-EXP-B", organization=self.org_b, projet=self.project_b,
            categorie="B", montant=90, created_by=self.manager_b,
        )
        Depense.objects.create(
            reference="LINEAGE-EXP-CROSS", organization=self.org_a, projet=self.project_b,
            categorie="Cross", montant=900, created_by=self.manager_a,
        )
        Don.objects.create(
            reference="LINEAGE-DON-A", organization=self.org_a, campagne=campaign_a,
            montant=5, created_by=self.manager_a,
        )
        Don.objects.create(
            reference="LINEAGE-DON-CROSS", organization=self.org_a, campagne=campaign_b,
            montant=800, created_by=self.manager_a,
        )

        dashboard = self.client.get("/api/reports/dashboard/")
        self.assertEqual(dashboard.status_code, 200)
        self.assertEqual(dashboard.data["resume"]["total_projets"], 1)
        self.assertEqual(dashboard.data["resume"]["total_campagnes"], 1)
        self.assertEqual(dashboard.data["resume"]["total_beneficiaires"], 1)
        self.assertEqual(dashboard.data["resume"]["total_depenses"], 10.0)
        self.assertEqual(dashboard.data["resume"]["budget_total"], 100.0)

        beneficiaries = self.client.get("/api/reports/beneficiaires/")
        self.assertEqual(beneficiaries.status_code, 200)
        self.assertEqual(beneficiaries.data["resume"]["total"], 1)
        self.assertEqual(beneficiaries.data["resume"]["total_femmes"], 1)
        self.assertEqual(
            [row["id"] for row in beneficiaries.data["top_vulnerables"]],
            [str(own_beneficiary.pk)],
        )

        financial = self.client.get("/api/reports/finance/")
        self.assertEqual(financial.status_code, 200)
        self.assertEqual(financial.data["resume"]["budget_total"], 100.0)
        self.assertEqual(financial.data["resume"]["total_depenses"], 10.0)
        self.assertEqual(financial.data["resume"]["total_dons"], 5.0)
        foreign_filter = self.client.get(
            "/api/reports/finance/", {"project_id": self.project_b.pk},
        )
        self.assertEqual(foreign_filter.status_code, 200)
        self.assertEqual(foreign_filter.data["resume"]["total_depenses"], 0.0)

        excel = self.client.get("/api/reports/export/excel/")
        self.assertEqual(excel.status_code, 200)
        import openpyxl
        from io import BytesIO
        workbook = openpyxl.load_workbook(BytesIO(excel.content), read_only=True)
        summary = list(workbook["Résumé"].values)
        self.assertEqual(summary[6][1], 1)  # beneficiary total
        self.assertEqual(summary[8][1], 10)  # expenses total
        project_rows = list(workbook["Projets"].values)
        self.assertEqual([row[1] for row in project_rows[1:]], ["Projet A"])
        beneficiary_rows = list(workbook["Bénéficiaires"].values)
        self.assertTrue(any(row[1] == "Beneficiary A" for row in beneficiary_rows))
        self.assertFalse(any(row[1] in {"Beneficiary B", "Foreign campaign link", "Foreign zone link"}
                             for row in beneficiary_rows))
        finance_rows = list(workbook["Finance"].values)
        self.assertEqual([row[0] for row in finance_rows[1:]], ["LINEAGE-EXP-A"])

        import base64
        import re
        import zlib
        pdf = self.client.get("/api/reports/export/pdf/")
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        self.assertTrue(pdf.content.startswith(b"%PDF-"))
        decoded_streams = []
        for match in re.finditer(rb"\bstream(?:\r\n|\n|\r)(.*?)\bendstream", pdf.content, re.DOTALL):
            stream = match.group(1).rstrip(b"\r\n")
            object_header = pdf.content[:match.start()].rsplit(b"obj", 1)[-1]
            try:
                if b"ASCII85Decode" in object_header:
                    stream = base64.a85decode(stream, adobe=True)
                if b"FlateDecode" in object_header:
                    stream = zlib.decompress(stream)
            except (ValueError, zlib.error):
                continue
            decoded_streams.append(stream)
        pdf_content = b"\n".join(decoded_streams)
        self.assertIn(b"Projet A", pdf_content)
        self.assertNotIn(b"Projet B", pdf_content)
        self.assertNotIn(b"ONG B", pdf_content)
        self.assertNotIn(b"Beneficiary B", pdf_content)
        self.assertNotIn(b"1 000", pdf_content)

        # The Celery PDF builder must use the same lineage checks as HTTP exports.
        from apps.reports import tasks
        from reportlab.platypus import Table
        from unittest.mock import patch
        captured = {}
        with patch("apps.reports.tasks.SimpleDocTemplate") as document:
            document.return_value.build.side_effect = lambda story: captured.update(story=story)
            tasks.build_pdf_report(self.org_a)
        summary_table = next(item for item in captured["story"] if isinstance(item, Table))
        self.assertEqual(summary_table._cellvalues[3][1], "1")
        self.assertEqual(summary_table._cellvalues[4][1], "100")
        self.assertEqual(summary_table._cellvalues[5][1], "10")

    def test_super_admin_cannot_use_organization_reports(self):
        super_admin = User.objects.create_user(
            email="reports-super@example.test", password="SafePass123!",
            role=User.Role.SUPER_ADMIN, status=User.Status.ACTIVE,
            is_active=True, is_staff=True,
        )
        client = APIClient()
        client.force_authenticate(super_admin)
        for path in (
            "/api/reports/dashboard/", "/api/reports/projets/",
            "/api/reports/campagnes/", "/api/reports/beneficiaires/",
            "/api/reports/finance/", "/api/reports/filters/",
            "/api/reports/export/pdf/", "/api/reports/export/excel/",
        ):
            with self.subTest(path=path):
                self.assertEqual(client.get(path).status_code, 403)


class PrivateGeneratedReportMediaTests(SimpleTestCase):
    def test_debug_media_fallback_cannot_serve_generated_reports(self):
        with TemporaryDirectory() as media_root:
            report_path = Path(media_root, "reports", "org-a.pdf")
            report_path.parent.mkdir()
            report_path.write_bytes(b"private report")
            with override_settings(
                DEBUG=True,
                ROOT_URLCONF="apps.reports.test_urls",
                MEDIA_ROOT=media_root,
            ):
                response = self.client.get("/media/reports/org-a.pdf")

        self.assertEqual(response.status_code, 404)
